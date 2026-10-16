"""多格式解析层：任何资料 -> **带位置信息的块**，再交给既有的切块器。

## 这一层解决什么

摄取链原本只认 `.md`（`iter_corpus_files` 里写死 `rglob("*.md")`），于是"把资料丢进来"
这件事对 PDF / Office / HTML 完全不成立 —— 不是检索质量差，而是**根本没进索引**，
且没有任何地方会说"你的文件被跳过了"。本模块把"读文件"从"切块"里拆出来：
解析器只负责 **抽出文字 + 记下这些字在哪**，切块仍然全部由 `chunker.py` 负责。

## 位置（`ParsedBlock.location`）是引用可信的前提

RAG 的引用如果不能指回原文的具体位置，就只是装饰。各格式的"位置"语义不同，
所以这里按格式分别定义，不做"统一成行号"这种会把语义抹平的归一：

| 格式 | 位置 | 切块粒度 |
| --- | --- | --- |
| `.pdf` | `p.7`（页码，1 起） | 按页 -> 页内再按段 |
| `.xlsx` / `.csv` | `Sheet1!A1:C20`（sheet 名 + 行区间） | 按行窗口，**每块重复表头行** |
| `.pptx` | `slide 3` | 按幻灯片（标题 + 要点） |
| `.docx` / `.html` / `.md` / `.txt` | 标题路径（h1–h6） | 交给 `chunk_document` 按标题切 |

`location` 会一路带到检索结果的引用里：存储层的字段是现成的
（`ChunkInput.breadcrumb` -> `ChunkHit.citation`），所以位置以"面包屑首段"的形式
进入引用（`[手册.pdf › p.7]`、`[台账.xlsx › Sheet1!A1:C20]`），不新增存储字段。

## 三条不能踩的线（都会毁掉现场演示）

1. **绝不静默丢字**：CSV/TXT 常见 GBK/GB18030。解码按 `utf-8-sig -> utf-8 -> gb18030`
   严格试，全失败时取"替换字符更少"的那个并用 `errors="replace"`，
   **把替换字符数写进报告**。`errors="ignore"` 会让文件"看起来解析成功"却缺字，
   检索不到时无从排查。
2. **扫描件必须点名**：PDF 抽出空文本 = 无文本层（不做 OCR）。这时**绝不建空块**，
   而是报"该文件无文本层（疑似扫描件），已跳过"并计数 —— "检索不到"必须能查出原因。
3. **单文件失败不影响整批**：`parse()` **不抛异常**，任何失败（不支持的格式、文件损坏、
   缺解析库）都变成 `ParsedDoc.error` 里的可读原因，由摄取层计入报告，本批继续。

解析库（pypdf / python-docx / openpyxl / python-pptx / beautifulsoup4）全部**在函数内
延迟 import**：它们已在 `pyproject.toml` 显式声明，但"缺一个库就让 `import kubedocs.ingest`
整体失败、连检索服务都起不来"是不可接受的降级路径。缺库时只坏这一种格式，且原因是可读的。
"""

from __future__ import annotations

import csv
import io
import os
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from kubedocs.chunker import DOC_SUFFIXES, parse_front_matter, strip_html

# 单文件大小上限。语料里混进一个几百 MB 的 PDF/PPT 会让一次摄取卡在解析上，
# 而且这类文件几乎都是扫描件（反正也抽不出文本）。超限的**不静默跳过**：计入报告。
MAX_FILE_BYTES = 64 * 1024 * 1024

NO_TEXT_LAYER_MSG = "该文件无文本层（疑似扫描件），已跳过"
EMPTY_CONTENT_MSG = "解析后没有可索引的文本，已跳过"

# 解码顺序：先 BOM 感知的 utf-8，再 gb18030（GBK 的超集，能覆盖繁体与部分生僻字）
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")
_UTF8_BOM = b"\xef\xbb\xbf"

# 表格一"块"的目标字符数。取值略小于默认 chunk_max_chars（1200），
# 这样一块通常刚好落成一个 chunk，表头重复一次的代价不会被二次切分抵消。
_TABLE_BLOCK_CHARS = 1000

# PDF 一行一段，段落目标长度取 400：既让下游 _split_long_text 有自然切点，
# 又不至于把一句话切成碎片。
_PARA_TARGET_CHARS = 400

# HTML 里纯噪声的标签：整块删掉（保留它们包裹的文字会让侧边栏/页脚污染检索）。
# `head` 也在内：`<title>`/`<meta>` 不是正文，而标题已经单独取出来了 ——
# 不删的话每篇 HTML 的第一行都是重复的标题。
_HTML_NOISE_TAGS = (
    "script",
    "style",
    "noscript",
    "template",
    "head",
    "nav",
    "footer",
    "form",
    "iframe",
    "svg",
)
_HTML_HEADING_TAG = re.compile(r"^h[1-6]$")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_DOCX_HEADING_RE = re.compile(r"(?:heading|标题)\s*([1-6])", re.IGNORECASE)


# --------------------------------------------------------------------------- 数据契约


@dataclass
class ParsedBlock:
    """解析器能确定的一个**结构单元**：一段文字 + 它在原文里的位置。

    `heading_path` 是块内的标题层级（`['指南', '安装']`）；`location` 是物理位置
    （`p.7` / `Sheet1!A1:C20` / `slide 3`）。标题型格式（md/html/docx）的
    `location` 留空 —— 它们的"位置"就是标题路径，由 `chunk_document` 从标题行派生，
    在这之前无法可靠地写死（正文里可能还没有标题）。
    """

    text: str
    location: str = ""
    heading_path: list[str] = field(default_factory=list)


@dataclass
class ParsedDoc:
    """一个文件的解析结果。`blocks` 为空 + `skip_reason` 非空 = 明确"跳过了，原因在此"。"""

    title: str = ""
    blocks: list[ParsedBlock] = field(default_factory=list)
    encoding: str = ""  # 文本类格式实际使用的编码（写进报告，便于解释乱码）
    scanned: bool = False  # PDF 抽不出任何文本 -> 疑似扫描件（不做 OCR）
    error: str = ""  # 无法解析的原因（不支持的格式 / 文件损坏 / 缺解析库）
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.error and not self.scanned and bool(self.blocks)

    @property
    def skip_reason(self) -> str:
        """没产出内容时的原因（空串 = 有内容）。调用方据此写报告。"""
        if self.error:
            return self.error
        if self.scanned:
            return NO_TEXT_LAYER_MSG
        if not self.blocks:
            return EMPTY_CONTENT_MSG
        return ""


@dataclass
class Discovery:
    """`discover_report` 的结果：可摄入的文件 + 被跳过的（超大）文件。"""

    files: list[Path] = field(default_factory=list)
    oversize: list[tuple[str, int]] = field(default_factory=list)  # (相对路径, 字节数)
    notes: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- 编码


def decode_text(raw: bytes) -> tuple[str, str, int]:
    """探测编码并解码，返回 `(文本, 编码名, 被替换的字符数)`。

    **绝不 `errors="ignore"`**：丢字是静默的，而"少了几行"在检索里表现为莫名其妙的
    漏召回。这里的策略是：能严格解码就严格解码（utf-8 -> gb18030）；
    两种编码都严格失败时，取"替换字符(U+FFFD)更少"的那个用 `replace` 救回来，
    并把替换数量返回给调用方计入报告 —— 宁可报告里写"3 个字节无法解码"，也不要假装成功。
    """
    for encoding in _TEXT_ENCODINGS:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        # 没有 BOM 就报 `utf-8`：报告里写 `utf-8-sig` 会让人以为文件带 BOM
        return text, ("utf-8" if encoding == "utf-8-sig" and not raw.startswith(_UTF8_BOM) else encoding), 0
    candidates = [
        (raw.decode("utf-8", errors="replace"), "utf-8"),
        (raw.decode("gb18030", errors="replace"), "gb18030"),
    ]
    text, encoding = min(candidates, key=lambda item: item[0].count("\ufffd"))
    return text, f"{encoding}(replace)", text.count("\ufffd")


def read_text(path: Path) -> tuple[str, str, list[str]]:
    """读文本文件：返回 `(文本, 编码名, 警告)`。警告里带上替换字符数，便于排查乱码。"""
    text, encoding, replaced = decode_text(path.read_bytes())
    warnings: list[str] = []
    if replaced:
        warnings.append(
            f"编码探测失败：utf-8 与 gb18030 都无法严格解码，已用 {encoding} 解码，"
            f"{replaced} 个字符被替换为 U+FFFD（**没有使用 errors=ignore 静默丢字**）"
        )
    return text, encoding, warnings


# --------------------------------------------------------------------------- 发现文件


def _rel_posix(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:  # root 是相对路径时的兜底（os.walk 的 dirpath 形态不一致）
        return path.name


def discover_report(root: Path, *, max_bytes: int = MAX_FILE_BYTES) -> Discovery:
    """递归发现受支持的文件，并记录被跳过的超大文件。

    跳过规则：
    - **隐藏目录/隐藏文件**（`.git`、`.cache`、`.venv`）整棵剪掉 —— 摄取时把
      `.git` 里的对象当语料搜出来是灾难；用 `os.walk` 原地剪枝而不是事后过滤，
      否则光遍历 `.git` 就能让一次摄取多跑几分钟；
    - 扩展名不在 `PARSERS` 里的文件不出现（不是错误，只是不归我们管）；
    - 超过 `max_bytes` 的文件**不静默跳过**：记进 `oversize` 由报告点名。

    返回顺序按相对路径排序：多机/多次运行结果一致，增量比对才有意义。
    """
    out = Discovery()
    if not root.exists():
        return out
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith("."))
        base = Path(dirpath)
        for name in sorted(filenames):
            if name.startswith("."):
                continue
            path = base / name
            if path.suffix.lower() not in PARSERS:
                continue
            try:
                if not path.is_file():
                    continue
                size = path.stat().st_size
            except OSError as exc:
                out.notes.append(f"无法读取文件属性，已跳过：{_rel_posix(path, root)}（{exc}）")
                continue
            if size > max_bytes:
                out.oversize.append((_rel_posix(path, root), size))
                continue
            out.files.append(path)
    out.files.sort(key=lambda p: _rel_posix(p, root))
    return out


def discover(root: Path, *, max_bytes: int = MAX_FILE_BYTES) -> list[Path]:
    """`discover_report` 的简化形态：只要文件列表（跳过原因见前者）。"""
    return discover_report(root, max_bytes=max_bytes).files


# --------------------------------------------------------------------------- 解析：文本型


def _title_from_markup(text: str) -> str:
    """front matter 的 title 优先，其次第一个标题行（标题型格式的通用取法）。"""
    meta, _ = parse_front_matter(text)
    title = (meta.get("title") or "").strip()
    if title:
        return title
    for line in text.splitlines():
        m = _MD_HEADING_RE.match(line)
        if m:
            return m.group(2).strip()
    return ""


def _parse_text(path: Path) -> ParsedDoc:
    """`.md` / `.txt`：只做解码与标题，**切块完全交给 `chunk_document`**。

    这里刻意不切分标题：K8s 语料是 1.7 万个 md、1.8 万个块，块边界与 content_hash
    是既有索引的增量基础（`ingest` 靠 content_hash 复用向量）。解析层再切一次标题
    就等于换一套切块规则 = 全量重切 + 引用面全变，所以 md/txt 的正文原样下传。
    """
    text, encoding, warnings = read_text(path)
    doc = ParsedDoc(title=_title_from_markup(text), encoding=encoding, warnings=warnings)
    if text.strip():
        doc.blocks.append(ParsedBlock(text=text))
    return doc


def _parse_html(path: Path) -> ParsedDoc:
    """`.html` / `.htm`：先剥标签，但把 h1–h6 变成 markdown 标题行**保留层级**。

    为什么不是"直接抽纯文本"：HTML 里的层级全在标签里，标签一剥就只剩一堆平铺的
    句子，切块时无从分段，引用也只能说"在这个网页里"。转成 `#` 标题后，
    `chunk_document` 就能给出与 md 语料完全一致的标题面包屑。
    `nav/footer/script/style` 是纯噪声（每个页面都一样），整块删掉；
    `<table>`（API 文档大量使用）交给 `chunker.strip_html` 变成可读行文本。
    """
    try:
        from bs4 import BeautifulSoup
        from bs4.element import NavigableString
    except ImportError as exc:  # pragma: no cover - 依赖已声明，装不上才走到这里
        return ParsedDoc(error=f"缺少解析库 beautifulsoup4（pip install beautifulsoup4）：{exc}")

    raw, encoding, warnings = read_text(path)
    soup = BeautifulSoup(raw, "html.parser")
    # 先取标题再删 `head`：`<title>` 是元数据不入正文，但它常常是唯一的标题来源
    h1 = soup.find("h1")
    title = ""
    if h1 is not None:
        title = h1.get_text(" ", strip=True)
    elif soup.title is not None:
        title = soup.title.get_text(" ", strip=True)
    for tag in soup.find_all(_HTML_NOISE_TAGS):
        tag.decompose()

    # 表格倒序处理：外层表格的字符串里已包含内层表格转换后的文本，正序处理会漏掉嵌套表
    for table in reversed(soup.find_all("table")):
        table.replace_with(NavigableString(f"\n\n{strip_html(str(table))}\n\n"))
    for heading in soup.find_all(_HTML_HEADING_TAG):
        level = int(heading.name[1])
        text = heading.get_text(" ", strip=True)
        heading.replace_with(NavigableString(f"\n\n{'#' * level} {text}\n\n"))

    body = _tidy_lines(soup.get_text("\n"))
    doc = ParsedDoc(title=title, encoding=encoding, warnings=warnings)
    if body:
        doc.blocks.append(ParsedBlock(text=body))
    return doc


def _tidy_lines(text: str) -> str:
    """规整行：去行尾空白、压掉连续空行，但**保留行首缩进**（代码块靠它可读）。"""
    out: list[str] = []
    for line in text.splitlines():
        line = line.rstrip()
        if not line and (not out or not out[-1]):
            continue
        out.append(line)
    return "\n".join(out).strip()


# --------------------------------------------------------------------------- 解析：docx


def _docx_heading_level(paragraph: Any) -> int:
    """段落的标题层级（0 = 正文）。

    要同时看 `style.name` 与 `style_id`：中文版 Word 的样式名是"标题 1"，而
    `style_id` 仍是 `Heading1`（反之，从某些工具导出的文档只有 name 是英文）。
    只看其中一个都会让整篇文档退化成"没有标题"。
    """
    style = paragraph.style
    if style is None:
        return 0
    for candidate in (getattr(style, "style_id", "") or "", style.name or ""):
        m = _DOCX_HEADING_RE.search(candidate)
        if m:
            return int(m.group(1))
        if candidate.strip().lower() in ("title", "标题"):
            return 1
    return 0


def _parse_docx(path: Path) -> ParsedDoc:
    """`.docx`：段落与表格**按文档顺序**还原成 markdown，标题样式变成 `#` 层级。

    用 `iter_inner_content()` 而不是 `document.paragraphs + document.tables`：
    后者会把表格全部挪到正文末尾，切块后的阅读顺序完全错乱。
    """
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        return ParsedDoc(error=f"缺少解析库 python-docx（pip install python-docx）：{exc}")

    document = Document(str(path))
    lines: list[str] = []
    title = ""
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            level = _docx_heading_level(item)
            if level and not title:
                title = text
            lines.append(f"{'#' * level} {text}" if level else text)
        elif isinstance(item, Table):
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            rendered = _render_rows(rows)
            if rendered:
                lines.append(rendered)

    body = "\n\n".join(lines).strip()
    doc = ParsedDoc(title=title)
    if body:
        doc.blocks.append(ParsedBlock(text=body))
    return doc


# --------------------------------------------------------------------------- 解析：表格（xlsx / csv）


def _cell(value: Any) -> str:
    """单元格 -> 检索友好的字符串：`None` 空串、整数值不带 `.0`、换行压平。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\n", " ").replace("|", "\\|").strip()


def _col_letter(width: int) -> str:
    """列号 -> 字母（1 -> A，27 -> AA）。位置写 `A1:C20` 就靠它。"""
    out = ""
    n = max(width, 1)
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _row_line(row: list[str], width: int) -> str:
    """一行 -> `a | b | c`（列数按表宽补齐，缺列不塌陷）。"""
    return " | ".join([*row, *([""] * (width - len(row)))])


def _render_rows(rows: list[list[str]]) -> str:
    """二维表 -> 每行一条 `a | b | c` 的文本（列数按最宽行补齐）。"""
    rows = [[_cell(c) for c in row] for row in rows]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    return "\n".join(_row_line(row, width) for row in rows)


def _table_blocks(rows: list[list[str]], sheet: str) -> list[ParsedBlock]:
    """二维表 -> 块列表，**每块都重复表头行**，位置带上 sheet 名与行区间。

    为什么表头必须重复：表里一行只有 `张三 42 男` 的块，永远匹配不到"年龄"这个词——
    而列名恰恰是提问里最常出现的词。重复表头的代价是每块多几十字符，
    换来的是"按列名检索得到数据行"。

    分块按字符数贪心（目标 `_TABLE_BLOCK_CHARS`）而不是固定行数：宽表和窄表的
    行长度可以差十倍，固定行数会让宽表的块远超嵌入模型上限、窄表的块全是碎片。
    """
    cleaned = [[_cell(c) for c in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return []
    width = max(len(row) for row in cleaned)
    header = cleaned[0]
    header_line = _row_line(header, width)
    body = cleaned[1:]
    if not body:
        return [ParsedBlock(text=header_line, location=f"{sheet}!A1:{_col_letter(width)}1")]

    blocks: list[ParsedBlock] = []
    i = 0
    while i < len(body):
        window = [header]
        size = len(header_line)
        while i < len(body):
            line = _row_line(body[i], width)
            if len(window) > 1 and size + len(line) + 1 > _TABLE_BLOCK_CHARS:
                break
            window.append(body[i])
            size += len(line) + 1
            i += 1
        first = i - len(window) + 3  # 表头占第 1 行，故数据行号 = 序号 + 1
        if first == 2:
            first = 1  # 首个窗口物理上就含表头行
        blocks.append(
            ParsedBlock(
                text="\n".join(_row_line(row, width) for row in window),
                location=f"{sheet}!A{first}:{_col_letter(width)}{i + 1}",
            )
        )
    return blocks


def _parse_xlsx(path: Path) -> ParsedDoc:
    """`.xlsx`：按 sheet 分块。`data_only=True` 读公式的**计算结果**而不是 `=SUM(...)`。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        return ParsedDoc(error=f"缺少解析库 openpyxl（pip install openpyxl）：{exc}")

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        blocks: list[ParsedBlock] = []
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            blocks.extend(_table_blocks(rows, sheet.title))
    finally:
        workbook.close()
    doc = ParsedDoc(title=path.stem, blocks=blocks)
    if not blocks:
        doc.warnings.append(EMPTY_CONTENT_MSG)
    return doc


def _parse_csv(path: Path) -> ParsedDoc:
    """`.csv`：与 xlsx 同一条表格路径（表头重复 + 行区间位置），编码同样探测。

    刻意复用 `_table_blocks`：CSV 与 XLSX 的检索需求完全一样（"按列名找到数据行"），
    分成两套只会让"表头别忘了重复"这条规则在其中一套里被忘掉。
    """
    text, encoding, warnings = read_text(path)
    try:
        rows = [list(row) for row in csv.reader(io.StringIO(text))]
    except csv.Error as exc:
        return ParsedDoc(encoding=encoding, error=f"CSV 解析失败：{exc}", warnings=warnings)
    blocks = _table_blocks(rows, path.stem)
    doc = ParsedDoc(title=path.stem, blocks=blocks, encoding=encoding, warnings=warnings)
    if not blocks:
        doc.warnings.append(EMPTY_CONTENT_MSG)
    return doc


# --------------------------------------------------------------------------- 解析：PDF


def _paragraphize(text: str) -> str:
    """PDF 抽出的文本"一行一段"，直接当一段会是一整页的巨型块。这里给它分段。

    先按空行切（PDF 内部有段落信息时最准），对仍超长的段再按行贪心装箱到
    ~400 字符 —— 下游 `_split_long_text` 才能找到自然切点，而不是把一整页塞进一个块。
    """
    out: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        if not para.strip():
            continue
        lines = [line.strip() for line in para.splitlines() if line.strip()]
        if len(para) <= _PARA_TARGET_CHARS or len(lines) <= 1:
            out.append(para.strip())
            continue
        cur: list[str] = []
        size = 0
        for line in lines:
            if cur and size + len(line) + 1 > _PARA_TARGET_CHARS:
                out.append("\n".join(cur))
                cur, size = [], 0
            cur.append(line)
            size += len(line) + 1
        if cur:
            out.append("\n".join(cur))
    return "\n\n".join(out)


def _parse_pdf(path: Path) -> ParsedDoc:
    """`.pdf`：一页一块（`location="p.N"`），页内再分段。

    **无文本层 = 明确跳过**：扫描件抽出来是空字符串。这时建一个空块是最坏的选择——
    "库里有这个文档但搜不到"从此无法解释；空向量还会触发 `_sanity_check` 直接中止整批。
    所以这里返回 `scanned=True`，由摄取层报"该文件无文本层（疑似扫描件），已跳过"。
    OCR 不在本模块职责内（要做也是接独立的 OCR 服务，而不是在这里假装抽到了文本）。
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        return ParsedDoc(error=f"缺少解析库 pypdf（pip install pypdf）：{exc}")

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # 加密 PDF 的原因千奇百怪，一律当"读不了"
            return ParsedDoc(error=f"PDF 已加密且无法用空口令解密：{exc}")

    title = ""
    try:
        meta = reader.metadata
        title = (getattr(meta, "title", "") or "").strip()
    except Exception:  # 元数据损坏不该让整篇 PDF 读不了
        title = ""

    blocks: list[ParsedBlock] = []
    empty_pages: list[int] = []
    warnings: list[str] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = _paragraphize(page.extract_text() or "")
        except Exception as exc:  # 单页失败不该毁掉整篇
            empty_pages.append(number)
            warnings.append(f"第 {number} 页抽取失败（{type(exc).__name__}: {exc}）")
            continue
        if not text:
            empty_pages.append(number)
            continue
        blocks.append(ParsedBlock(text=text, location=f"p.{number}"))

    if not blocks:
        return ParsedDoc(title=title, scanned=True, warnings=[*warnings, f"共 {len(empty_pages)} 页无文本层"])
    if empty_pages:
        shown = "、".join(str(p) for p in empty_pages[:10])
        more = f"（另 {len(empty_pages) - 10} 页未列出）" if len(empty_pages) > 10 else ""
        warnings.append(f"第 {shown} 页无文本层，已跳过{more}")
    return ParsedDoc(title=title, blocks=blocks, warnings=warnings)


# --------------------------------------------------------------------------- 解析：pptx


def _parse_pptx(path: Path) -> ParsedDoc:
    """`.pptx`：一页幻灯片一块（`location="slide N"`），取标题 + 正文要点 + 备注。

    备注（notes）也收进来：很多汇报把真正的说明放在备注里，只抽形状文本会漏掉一半内容。
    """
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        return ParsedDoc(error=f"缺少解析库 python-pptx（pip install python-pptx）：{exc}")

    presentation = Presentation(str(path))
    core_title = ""
    try:
        core_title = (presentation.core_properties.title or "").strip()
    except Exception:  # 元数据缺字段在真实文件里很常见
        core_title = ""

    blocks: list[ParsedBlock] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_title = ""
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            slide_title = next(
                (line.strip() for line in title_shape.text_frame.text.splitlines() if line.strip()), ""
            )

        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                rendered = _render_rows(rows)
                if rendered:
                    lines.append(rendered)
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            if title_shape is not None and shape == title_shape:
                continue  # 标题单独放在最前面，不在要点里重复
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                lines.append(f"{'  ' * max(paragraph.level, 0)}- {text}")

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        parts = [p for p in (slide_title, "\n".join(lines), notes) if p]
        text = "\n\n".join(parts)
        if not text:
            continue
        blocks.append(
            ParsedBlock(
                text=text, location=f"slide {index}", heading_path=[slide_title] if slide_title else []
            )
        )

    doc = ParsedDoc(title=core_title or path.stem, blocks=blocks)
    if not blocks:
        doc.warnings.append(EMPTY_CONTENT_MSG)
    return doc


# --------------------------------------------------------------------------- 注册表与入口

Parser = Callable[[Path], ParsedDoc]

# 后缀 -> 解析器。key 必须**恰好覆盖** `chunker.DOC_SUFFIXES`
# （tests/unit/test_parse.py 有一条断言钉住这个不变式：两份清单只允许一处定义、
#  另一处对齐，否则会出现"发现得了但解析不了"的文件）。
PARSERS: dict[str, Parser] = {
    ".md": _parse_text,
    ".txt": _parse_text,
    ".csv": _parse_csv,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".pptx": _parse_pptx,
}


def parse(path: Path) -> ParsedDoc:
    """解析一个文件。**不抛异常**：任何失败都变成 `ParsedDoc.error` 里的可读原因。

    摄取是批处理，"一个坏文件炸掉整批"意味着剩下的几百个文件一个都没进索引，
    而用户看到的只是一个栈。所以这里把异常收口：不支持的格式给出受支持清单，
    解析器的任何异常都带上类型与消息（文件由调用方补进报告）。
    """
    suffix = path.suffix.lower()
    parser = PARSERS.get(suffix)
    if parser is None:
        supported = "、".join(sorted(PARSERS))
        return ParsedDoc(error=f"不支持的格式：{suffix or '（无扩展名）'}；受支持的格式：{supported}")
    try:
        doc = parser(path)
    except Exception as exc:  # 单文件失败不许影响整批，原因原样上报
        return ParsedDoc(error=f"{type(exc).__name__}: {exc}")
    if not doc.title:
        doc.title = path.stem
    return doc


def supported_suffixes() -> tuple[str, ...]:
    """受支持的后缀（与 `chunker.DOC_SUFFIXES` 同源，供文档/CLI 展示）。"""
    return DOC_SUFFIXES
