"""解析层契约测试：每种格式都真的抽出"文字 + 位置"，编码不丢字，坏文件不炸整批。

全部离线：样本一律用 `tmp_path` 现场生成 —— docx/xlsx/pptx 用各自的库写，PDF 用手写的
最小 PDF 字节串（含 `BT/Tj/ET` 与正确 xref，不额外引依赖），不碰 `data/` 里的真实语料、
不联网、不调模型。
"""

from __future__ import annotations

import os
import re
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from kubedocs.chunker import DOC_SUFFIXES, chunk_blocks, chunk_document
from kubedocs.ingest import ingest
from kubedocs.parse import (
    MAX_FILE_BYTES,
    NO_TEXT_LAYER_MSG,
    PARSERS,
    ParsedBlock,
    discover,
    discover_report,
    parse,
)
from kubedocs.store.base import ChunkHit
from kubedocs.store.sqlite_store import SqliteStore

# 标题型格式的正文要**长于 min_chars(80)**，否则会被"碎片并入前一块"的规则合并成一块，
# 断言面包屑时就看不出层级了。
LONG = "就诊须知：先取号再候诊，携带身份证与医保卡，过号请重新取号。" * 3
LONG2 = "环境要求：Python 3.10 以上，内存 8GB 以上，否则嵌入模型加载会失败。" * 3
HEADER = "医生 | 科室 | 出诊时间"
ROW = "心内科 | 周一上午 08:00-12:00 门诊三楼 3 号诊室"


# --------------------------------------------------------------------------- 样本构造


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)").replace("\n", r"\n")


def _pdf_bytes(pages: list[str]) -> bytes:
    """手写最小合法 PDF：每页一个内容流（`BT /F1 12 Tf ... Tj ET`）+ 正确的 xref。

    手写而不是引第三方库生成：测试要能证明"我们解析的是真实 PDF 结构"，
    而不是"某个库写出来的库再读回来"。换行用 `\\n` 转义（PDF 字面串里的换行），
    pypdf 会还原成真正的换行。
    """
    count = len(pages)
    font_obj = 3 + 2 * count
    objs: dict[int, bytes] = {}
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(count))
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode()
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({_pdf_escape(text)}) Tj ET".encode()
        objs[3 + 2 * i] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> /Contents {4 + 2 * i} 0 R >>"
        ).encode()
        objs[4 + 2 * i] = b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
    objs[font_obj] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n"
    xref_at = len(out)
    size = max(objs) + 1
    out += b"xref\n0 %d\n" % size + b"0000000000 65535 f \n"
    for num in range(1, size):
        out += b"%010d 00000 n \n" % offsets.get(num, 0)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref_at)
    return bytes(out)


def _write_xlsx(path: Path, rows: int = 80) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "门诊排班"
    sheet.append(["医生", "科室", "出诊时间"])
    for i in range(rows):
        sheet.append([f"医生{i}", "心内科", "周一上午 08:00-12:00 门诊三楼 3 号诊室"])
    prices = workbook.create_sheet("价目表")
    prices.append(["项目", "价格"])
    prices.append(["血常规", 35])
    workbook.save(str(path))


def _write_pptx(path: Path) -> None:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    slides = [
        ("项目背景", ["任何格式都要能被检索", "位置信息必须进引用"]),
        ("结论", ["先接解析层，再谈质量"]),
    ]
    for title, bullets in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0]
        for extra in bullets[1:]:
            frame.add_paragraph().text = extra
    presentation.save(str(path))


def _first_hit(store: SqliteStore, query: str) -> ChunkHit:
    ranked = store.keyword_search(query, 5)
    assert ranked, f"关键词 {query!r} 一个块都没检索到"
    return store.get_chunks([ranked[0][0]])[ranked[0][0]]


# --------------------------------------------------------------------------- 注册表


def test_注册表恰好覆盖受支持后缀():
    """两份清单只允许一处定义、另一处对齐：否则会出现"发现得了但解析不了"的文件。"""
    assert set(PARSERS) == set(DOC_SUFFIXES)


def test_不支持的格式返回可读原因而不抛异常(tmp_path: Path):
    path = tmp_path / "note.bin"
    path.write_bytes(b"\x00\x01\x02")
    doc = parse(path)  # 不 raise 是契约的一部分
    assert doc.blocks == []
    assert ".bin" in doc.error and "不支持" in doc.error
    assert doc.skip_reason == doc.error, "读不了的原因必须能直接被报告引用"
    assert doc.ok is False


@pytest.mark.parametrize("name", ["broken.pdf", "broken.xlsx", "broken.docx", "broken.pptx"])
def test_损坏的文件变成可读原因而不是异常(tmp_path: Path, name: str):
    path = tmp_path / name
    path.write_bytes(b"PK\x03\x04 this is not a real document")
    doc = parse(path)
    assert doc.error, f"{name} 应当给出可读失败原因"
    assert doc.blocks == []
    assert "Traceback" not in doc.error


# --------------------------------------------------------------------------- md / txt


def test_md_正文原样下传且位置等于标题路径(tmp_path: Path):
    path = tmp_path / "guide.md"
    path.write_text(
        f"---\ntitle: 门诊手册\n---\n\n# 门诊手册\n\n{LONG}\n\n## 挂号流程\n\n{LONG2}\n",
        encoding="utf-8",
    )
    doc = parse(path)
    assert doc.title == "门诊手册"
    assert len(doc.blocks) == 1, "md 的正文整篇下传，标题层级由既有 chunk_document 派生"

    chunks = chunk_document("guide.md", doc.blocks[0].text)
    assert [c.breadcrumb for c in chunks] == [["门诊手册"], ["门诊手册", "挂号流程"]]
    # 块的 location 与面包屑一致（都是标题路径）—— 引用里因此不会重复出现位置
    assert [c.location for c in chunks] == ["门诊手册", "门诊手册 › 挂号流程"]
    assert all(c.source_url == "https://kubernetes.io/docs/guide/" for c in chunks)


def test_txt_纯文本也能解析(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("复诊提醒：带好病历本与身份证。\n", encoding="utf-8")
    doc = parse(path)
    assert doc.encoding == "utf-8"  # 没有 BOM 就报 utf-8，不报 utf-8-sig
    assert doc.blocks[0].text.strip() == "复诊提醒：带好病历本与身份证。"
    assert doc.title == "notes", "没有标题时回退为文件名"


# --------------------------------------------------------------------------- 编码


def test_gbk_txt_解码不丢字(tmp_path: Path):
    """GBK/GB18030 是中文语料的常态：必须解对，且一个字都不能少。

    样本里放 `𠀀`(U+20000)：GBK **编不出来**，只有 gb18030 可以 —— 这条断言同时证明
    回退到的是 gb18030 而不只是 GBK。
    """
    text = "复诊提醒：带好病历本与身份证。𠀀\n"
    with pytest.raises(UnicodeEncodeError):
        text.encode("gbk")  # 证明这个样本确实超出 GBK
    path = tmp_path / "gbk.txt"
    path.write_bytes(text.encode("gb18030"))

    doc = parse(path)
    assert doc.encoding == "gb18030"
    assert doc.blocks[0].text == text
    assert "𠀀" in doc.blocks[0].text
    assert "\ufffd" not in doc.blocks[0].text


def test_极度损坏的编码_统计替换字符而不是静默丢字(tmp_path: Path):
    """`errors="ignore"` 会让文件"看起来解析成功"却缺字 —— 这里必须留下痕迹。"""
    path = tmp_path / "bad.txt"
    path.write_bytes(b"abc\x81\xff\xff\xff hello")
    doc = parse(path)
    assert doc.encoding.endswith("(replace)")
    assert any("U+FFFD" in w and "ignore" in w for w in doc.warnings)
    assert doc.blocks, "不能因为编码不完美就整篇丢掉"


# --------------------------------------------------------------------------- HTML


def test_html_剥标签保留标题层级并丢掉_nav_footer(tmp_path: Path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>安装指南</title></head><body>"
        "<nav>首页 关于 联系我们</nav>"
        "<h1>安装指南</h1>"
        f"<p>{LONG2}</p>"
        "<h2>环境要求</h2>"
        f"<p>{LONG}</p>"
        "<table><tr><th>名称</th><th>版本</th></tr><tr><td>python</td><td>3.10</td></tr></table>"
        "<footer>版权所有 2026</footer></body></html>",
        encoding="utf-8",
    )
    doc = parse(path)
    body = doc.blocks[0].text
    assert doc.title == "安装指南"
    assert "# 安装指南" in body and "## 环境要求" in body, "h1–h6 必须保留成标题层级"
    assert "首页" not in body and "联系我们" not in body and "版权所有" not in body
    assert "Python 3.10 以上" in body, "正文不能一起被丢掉"
    assert "名称 | 版本" in body and "python | 3.10" in body, "表格要变成可读行文本"

    chunks = chunk_document("page.html", body)
    assert [c.breadcrumb for c in chunks] == [["安装指南"], ["安装指南", "环境要求"]]


# --------------------------------------------------------------------------- docx


def test_docx_标题样式与表格都进正文(tmp_path: Path):
    path = tmp_path / "sop.docx"
    document = Document()
    document.add_heading("检验科SOP", level=1)
    document.add_paragraph(LONG)
    document.add_heading("标本接收", level=2)
    document.add_paragraph(LONG2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "步骤"
    table.cell(0, 1).text = "责任人"
    table.cell(1, 0).text = "接收标本"
    table.cell(1, 1).text = "护士"
    document.save(str(path))

    doc = parse(path)
    body = doc.blocks[0].text
    assert doc.title == "检验科SOP"
    assert "# 检验科SOP" in body and "## 标本接收" in body
    assert "步骤 | 责任人" in body and "接收标本 | 护士" in body, "docx 表格里的字也必须能被检索到"

    chunks = chunk_document("sop.docx", body)
    assert [c.breadcrumb for c in chunks] == [["检验科SOP"], ["检验科SOP", "标本接收"]]


# --------------------------------------------------------------------------- xlsx / csv


def test_xlsx_多sheet且每块重复表头(tmp_path: Path):
    path = tmp_path / "schedule.xlsx"
    _write_xlsx(path, rows=80)

    doc = parse(path)
    sheet1 = [b for b in doc.blocks if b.location.startswith("门诊排班!")]
    prices = [b for b in doc.blocks if b.location.startswith("价目表!")]
    assert len(sheet1) >= 3, "80 行长表必须切成多块，否则单块远超嵌入长度"
    assert len(prices) == 1

    for block in sheet1:
        assert re.fullmatch(r"门诊排班!A\d+:C\d+", block.location), block.location
        assert block.text.splitlines()[0] == HEADER, "**每一块**都必须重复表头行"
        assert ROW in block.text
    assert sheet1[0].location.startswith("门诊排班!A1:")
    # 各块从不同的数据行开始（不是同一批行重复切）
    assert len({b.text.splitlines()[1] for b in sheet1}) == len(sheet1)
    assert "医生0" in sheet1[0].text and "医生79" in sheet1[-1].text
    assert prices[0].location == "价目表!A1:B2"
    assert "血常规 | 35" in prices[0].text, "整数不该渲染成 35.0"

    chunks = chunk_blocks("schedule.xlsx", doc.blocks, title=doc.title)
    assert all(c.location for c in chunks)
    assert chunks[0].text.startswith("[schedule › 门诊排班!A1:")
    assert chunks[0].source_url == "https://kubernetes.io/docs/schedule.xlsx"


def test_csv_按表格解析且编码可探测(tmp_path: Path):
    path = tmp_path / "价目表.csv"
    path.write_bytes(("项目,价格\n血常规,35\n尿常规,20\n").encode("gb18030"))
    doc = parse(path)
    assert doc.encoding == "gb18030"
    assert len(doc.blocks) == 1
    assert doc.blocks[0].location == "价目表!A1:B3"
    assert doc.blocks[0].text.splitlines()[0] == "项目 | 价格"
    assert "血常规 | 35" in doc.blocks[0].text


# --------------------------------------------------------------------------- pptx


def test_pptx_按页切分且标题进引用(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    _write_pptx(path)

    doc = parse(path)
    assert [b.location for b in doc.blocks] == ["slide 1", "slide 2"]
    assert doc.blocks[0].heading_path == ["项目背景"]

    chunks = chunk_blocks("deck.pptx", doc.blocks, title=doc.title)
    assert [c.location for c in chunks] == ["slide 1", "slide 2"]
    assert chunks[0].breadcrumb == ["slide 1", "项目背景"]
    assert "任何格式都要能被检索" in chunks[0].text, "要点不能丢"
    assert chunks[1].text.startswith("[deck › slide 2 › 结论]"), chunks[1].text[:40]
    assert all(c.source_url == "https://kubernetes.io/docs/deck.pptx" for c in chunks)


# --------------------------------------------------------------------------- PDF


def test_pdf_手写最小PDF抽到文本与页码(tmp_path: Path):
    path = tmp_path / "handbook.pdf"
    path.write_bytes(
        _pdf_bytes(
            [
                "Clinic handbook page one about registration and triage.",
                "Page two: bring your ID card and insurance card.",
            ]
        )
    )
    doc = parse(path)
    assert doc.error == "" and doc.scanned is False
    assert [b.location for b in doc.blocks] == ["p.1", "p.2"]
    assert "registration" in doc.blocks[0].text
    assert "insurance card" in doc.blocks[1].text

    chunks = chunk_blocks("handbook.pdf", doc.blocks, title=doc.title)
    assert chunks[0].text.startswith("[handbook › p.1]")
    assert chunks[1].breadcrumb == ["p.2"]


def test_pdf_无文本层_报告为扫描件且绝不建空块(tmp_path: Path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(_pdf_bytes([""]))
    doc = parse(path)
    assert doc.scanned is True
    assert doc.blocks == [], "空块会让'库里有但搜不到'无从排查，必须不产出块"
    assert doc.error == "", "这不是异常，是明确跳过"
    assert doc.skip_reason == NO_TEXT_LAYER_MSG
    assert "无文本层" in doc.skip_reason and "扫描件" in doc.skip_reason and "已跳过" in doc.skip_reason


def test_pdf_长页拆成多块但位置仍在同一页(tmp_path: Path):
    long_page = "\n".join(f"Line {i} of the long page content" for i in range(200))
    path = tmp_path / "long.pdf"
    path.write_bytes(_pdf_bytes([long_page, "second page"]))
    doc = parse(path)
    assert [b.location for b in doc.blocks] == ["p.1", "p.2"]

    chunks = chunk_blocks("long.pdf", doc.blocks, max_chars=600, overlap=50, min_chars=10)
    page_one = [c for c in chunks if c.location == "p.1"]
    assert len(page_one) >= 2, "一页 6000 字符必须切成多块"
    assert all(c.breadcrumb == ["p.1"] for c in page_one), "同页切出的块位置不得漂移"
    assert all("p.1" in c.text.splitlines()[0] for c in page_one)


def test_容器格式_不跨位置合并碎片():
    """跨位置合并会把下一页/下一张幻灯片的文字挂到上一页的引用下。"""
    adjacent = chunk_blocks(
        "d.pdf",
        [ParsedBlock(text="A" * 200, location="p.1"), ParsedBlock(text="很短的一段", location="p.2")],
        min_chars=80,
    )
    assert [c.location for c in adjacent] == ["p.1", "p.2"]

    same_page = chunk_blocks(
        "d.pdf",
        [ParsedBlock(text="A" * 200, location="p.1"), ParsedBlock(text="很短的一段", location="p.1")],
        min_chars=80,
    )
    assert len(same_page) == 1, "同一位置内的碎片照旧合并"


# --------------------------------------------------------------------------- discover


def test_discover_嵌套隐藏目录超大文件与不支持后缀(tmp_path: Path):
    root = tmp_path / "corpus"
    (root / "sub").mkdir(parents=True)
    (root / ".git").mkdir()
    (root / ".cache").mkdir()
    (root / "b.md").write_text("b", encoding="utf-8")
    (root / "a.md").write_text("a", encoding="utf-8")
    (root / "sub" / "c.txt").write_text("c", encoding="utf-8")
    (root / "sub" / "skip.bin").write_bytes(b"x")
    (root / ".git" / "d.md").write_text("d", encoding="utf-8")
    (root / ".cache" / "e.md").write_text("e", encoding="utf-8")
    (root / ".hidden.md").write_text("f", encoding="utf-8")
    (root / "big.pdf").write_bytes(b"x" * 4096)

    assert [p.relative_to(root).as_posix() for p in discover(root)] == [
        "a.md",
        "b.md",
        "big.pdf",
        "sub/c.txt",
    ], "排序稳定 + 隐藏目录/隐藏文件/不支持后缀都被跳过"

    report = discover_report(root, max_bytes=1024)
    assert [rel for rel, _ in report.oversize] == ["big.pdf"]
    assert [p.name for p in report.files] == ["a.md", "b.md", "c.txt"]
    assert MAX_FILE_BYTES == 64 * 1024 * 1024, "默认上限写进报告，改了要同步改测试"


# --------------------------------------------------------------------------- 端到端


def test_ingest_混合格式_位置进引用_坏文件不炸整批(settings, store, embedder):
    corpus = settings.corpus_path
    (corpus / "guide.md").write_text(
        f"---\ntitle: 门诊手册\n---\n\n# 门诊手册\n\n{LONG}\n\n## 挂号流程\n\n{LONG2}\n",
        encoding="utf-8",
    )
    (corpus / "handbook.pdf").write_bytes(
        _pdf_bytes(["Clinic handbook page one about registration.", "Page two: bring your ID card."])
    )
    (corpus / "scan.pdf").write_bytes(_pdf_bytes([""]))
    (corpus / "broken.xlsx").write_bytes(b"PK\x03\x04 this is not a real workbook")
    _write_xlsx(corpus / "schedule.xlsx", rows=6)
    huge = corpus / "huge.pdf"
    with open(huge, "wb") as fh:
        fh.write(b"%PDF-1.4\n")
    os.truncate(huge, MAX_FILE_BYTES + 1)  # 稀疏文件：不真的写 64MB

    report = ingest(store, embedder, settings)

    assert report.files_scanned == 5, "超大文件在发现阶段就出局，不计入扫描"
    assert report.files_oversize == 1
    assert report.files_no_text == 1
    assert report.files_failed == 1
    assert any("无文本层" in n and "扫描件" in n and "scan.pdf" in n for n in report.notes)
    assert any("解析失败" in n and "broken.xlsx" in n for n in report.notes)
    assert any("上限" in n and "huge.pdf" in n for n in report.notes)
    assert any("文本编码" in n for n in report.notes), "报告里必须写清用了哪种编码"
    assert "无文本层跳过 1 个" in report.summary() and "解析失败 1 个" in report.summary()

    # 坏文件不登记（可能只是暂时读不了，下次要重试）；扫描件登记（固有属性，不重复跑）
    registered = set(store.existing_doc_hashes())
    assert "broken.xlsx" not in registered
    assert "scan.pdf" in registered
    assert store.stats().n_docs == 4

    # 位置一路进到引用：PDF 是页码、xlsx 是 sheet + 行区间
    assert _first_hit(store, "registration").breadcrumb == "p.1"
    assert _first_hit(store, "registration").citation == "handbook › p.1"
    schedule_hit = _first_hit(store, "心内科")
    assert schedule_hit.breadcrumb.startswith("门诊排班!A"), schedule_hit.breadcrumb
    assert HEADER in schedule_hit.body, "表头行必须跟着数据行一起被检索到"
    assert schedule_hit.source_url.endswith("/schedule.xlsx")
    # md 的位置就是标题路径：不得被重复拼一遍
    assert _first_hit(store, "挂号流程").breadcrumb == "门诊手册 › 挂号流程"
    assert store.verify_fts() == "ok"


def test_ingest_第二次跑全部跳过_包括扫描件(settings, store, embedder):
    """扫描件与空文档登记后不该每次白跑；坏文件每次都要重新报（它可能已经修好了）。"""
    corpus = settings.corpus_path
    (corpus / "guide.md").write_text(f"# 手册\n\n{LONG}\n", encoding="utf-8")
    (corpus / "scan.pdf").write_bytes(_pdf_bytes([""]))

    first = ingest(store, embedder, settings)
    assert first.files_no_text == 1
    second = ingest(store, embedder, settings)
    assert second.files_changed == 0, f"扫描件/空文档也算处理过：{second.notes}"
    assert second.files_no_text == 0
    assert second.files_unchanged == 2


def test_大量坏文件时报告只计数不刷屏(settings, store, embedder):
    """计数永远准；细节行有上限 —— 5000 个坏文件不该变成 5000 行日志。"""
    corpus = settings.corpus_path
    (corpus / "good.md").write_text(f"# 手册\n\n{LONG}\n", encoding="utf-8")
    for i in range(60):
        (corpus / f"bad{i:02d}.xlsx").write_bytes(b"PK\x03\x04 not a workbook")

    report = ingest(store, embedder, settings)

    assert report.files_failed == 60
    assert report.notes_dropped > 0
    assert any("未逐条列出" in n for n in report.notes)
    assert store.stats().n_docs == 1, "好文件照旧入库，坏文件不登记"
    assert store.keyword_search("手册", 5), "整批没有被坏文件带崩"
